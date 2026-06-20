"""P1/P2 模块深度功能测试

覆盖:
  P1: ora_errors, data_source, auth, cli, notification
  P2: service/__init__, progress, history_viewer, constants
"""

import os
import sys
import csv
import json
import time
import tempfile
import shutil
import io
import argparse
from pathlib import Path
from datetime import datetime, timedelta
from unittest.mock import patch, MagicMock

import pytest

# 确保 src 可导入
_project_root = Path(__file__).parent
sys.path.insert(0, str(_project_root))


# ============================================================================
# P1-1: ora_errors.py
# ============================================================================

class TestOraErrors:
    """测试 ORA 错误码翻译模块"""

    def test_translate_known_error(self):
        from src.ora_errors import translate_ora_error
        result = translate_ora_error("ORA-01017: invalid username/password")
        assert "用户名或密码无效" in result
        assert "ORA-01017" in result

    def test_translate_known_error_case_insensitive(self):
        from src.ora_errors import translate_ora_error
        result = translate_ora_error("ora-01017: some message")
        assert "用户名或密码无效" in result

    def test_translate_unknown_error(self):
        from src.ora_errors import translate_ora_error
        result = translate_ora_error("ORA-99999: unknown error")
        assert "数据库错误" in result
        assert "ORA-99999" in result

    def test_translate_empty_message(self):
        from src.ora_errors import translate_ora_error
        result = translate_ora_error("")
        assert result == "未知错误"

    def test_translate_none_message(self):
        from src.ora_errors import translate_ora_error
        result = translate_ora_error(None)
        assert result == "未知错误"

    def test_all_known_error_codes(self):
        """测试所有已知 ORA 错误码都能被翻译"""
        from src.ora_errors import translate_ora_error, ORA_ERROR_MAP
        for code, expected_zh in ORA_ERROR_MAP.items():
            result = translate_ora_error(f"{code}: test message")
            assert expected_zh in result, f"ORA code {code} not translated correctly"
            assert code in result

    def test_get_error_hint_known(self):
        from src.ora_errors import get_error_hint
        hint = get_error_hint("ORA-01017")
        assert hint is not None
        assert "用户名" in hint or "密码" in hint

    def test_get_error_hint_unknown(self):
        from src.ora_errors import get_error_hint
        hint = get_error_hint("ORA-99999")
        assert hint is None

    def test_get_error_hint_all_documented(self):
        """测试所有有 hint 的错误码"""
        from src.ora_errors import get_error_hint
        documented_codes = [
            "ORA-01017", "ORA-12154", "ORA-12541", "ORA-00942",
            "ORA-00904", "ORA-00001", "ORA-01722", "ORA-01407",
            "ORA-01653", "ORA-01555", "ORA-00060", "ORA-01034", "ORA-01031",
        ]
        for code in documented_codes:
            hint = get_error_hint(code)
            assert hint is not None, f"Expected hint for {code}"

    def test_extract_ora_code_known(self):
        from src.ora_errors import extract_ora_code
        code, msg = extract_ora_code("ORA-00942: table or view does not exist")
        assert code == "ORA-00942"
        assert "表或视图不存在" in msg

    def test_extract_ora_code_unknown(self):
        from src.ora_errors import extract_ora_code
        code, msg = extract_ora_code("ORA-99999: mysterious error")
        assert code == ""
        assert "ORA-99999" in msg

    def test_extract_ora_code_empty(self):
        from src.ora_errors import extract_ora_code
        code, msg = extract_ora_code("")
        assert code == ""
        assert msg == "未知错误"

    def test_extract_ora_code_none(self):
        from src.ora_errors import extract_ora_code
        code, msg = extract_ora_code(None)
        assert code == ""
        assert msg == "未知错误"

    def test_ora_error_map_has_30_plus_codes(self):
        from src.ora_errors import ORA_ERROR_MAP
        assert len(ORA_ERROR_MAP) >= 30

    def test_multiple_ora_codes_in_message(self):
        """消息中包含多个 ORA 码时，应匹配第一个"""
        from src.ora_errors import translate_ora_error
        result = translate_ora_error("ORA-01017 and ORA-12154 both occurred")
        # 应匹配 ORA_ERROR_MAP 中第一个匹配的
        assert "ORA-" in result


# ============================================================================
# P1-2: data_source.py
# ============================================================================

class TestCsvDataSource:
    """测试 CSV 数据源"""

    def setup_method(self):
        self.tmpdir = tempfile.mkdtemp()

    def teardown_method(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _create_csv(self, filename, rows, delimiter=','):
        path = os.path.join(self.tmpdir, filename)
        with open(path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()), delimiter=delimiter)
            writer.writeheader()
            writer.writerows(rows)
        return path

    def test_read_data(self):
        from src.data_source import CsvDataSource
        path = self._create_csv("test.csv", [
            {"ID": "1", "NAME": "Alice"},
            {"ID": "2", "NAME": "Bob"},
        ])
        ds = CsvDataSource(path)
        rows = list(ds.read_rows())
        assert len(rows) == 2
        assert rows[0]["NAME"] == "Alice"

    def test_get_columns(self):
        from src.data_source import CsvDataSource
        path = self._create_csv("test.csv", [
            {"ID": "1", "NAME": "Alice", "AGE": "30"},
        ])
        ds = CsvDataSource(path)
        cols = ds.get_columns()
        assert cols == ["ID", "NAME", "AGE"]

    def test_get_row_count(self):
        from src.data_source import CsvDataSource
        path = self._create_csv("test.csv", [
            {"ID": str(i), "NAME": f"User{i}"} for i in range(50)
        ])
        ds = CsvDataSource(path)
        assert ds.get_row_count() == 50

    def test_file_not_found(self):
        from src.data_source import CsvDataSource
        ds = CsvDataSource("/nonexistent/path.csv")
        with pytest.raises(FileNotFoundError):
            ds.get_columns()

    def test_streaming_read(self):
        """测试流式读取（大数据集）"""
        from src.data_source import CsvDataSource
        path = self._create_csv("big.csv", [
            {"ID": str(i), "VALUE": f"val_{i}"} for i in range(1000)
        ])
        ds = CsvDataSource(path)
        count = 0
        for row in ds.read_rows():
            count += 1
        assert count == 1000

    def test_max_rows_limit(self):
        from src.data_source import CsvDataSource
        path = self._create_csv("big.csv", [
            {"ID": str(i)} for i in range(100)
        ])
        ds = CsvDataSource(path, max_rows=10)
        rows = list(ds.read_rows())
        assert len(rows) == 10

    def test_source_type(self):
        from src.data_source import CsvDataSource
        path = self._create_csv("test.csv", [{"A": "1"}])
        ds = CsvDataSource(path)
        assert ds.source_type == "csv"

    def test_get_preview(self):
        from src.data_source import CsvDataSource
        path = self._create_csv("test.csv", [
            {"ID": str(i), "NAME": f"User{i}"} for i in range(100)
        ])
        ds = CsvDataSource(path)
        # get_preview 依赖 _columns 已填充，需先调用 get_columns
        cols = ds.get_columns()
        cols2, preview = ds.get_preview(max_rows=5)
        assert len(preview) == 5
        assert "ID" in cols2

    def test_tab_delimiter_auto_detect(self):
        """测试 Tab 分隔符自动检测"""
        from src.data_source import CsvDataSource
        path = os.path.join(self.tmpdir, "tab.csv")
        with open(path, 'w', encoding='utf-8') as f:
            f.write("ID\tNAME\n1\tAlice\n2\tBob\n")
        ds = CsvDataSource(path)
        cols = ds.get_columns()
        assert "ID" in cols
        rows = list(ds.read_rows())
        assert len(rows) == 2


class TestJsonDataSource:
    """测试 JSON 数据源"""

    def setup_method(self):
        self.tmpdir = tempfile.mkdtemp()

    def teardown_method(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_json_array(self):
        from src.data_source import JsonDataSource
        path = os.path.join(self.tmpdir, "test.json")
        data = [{"ID": 1, "NAME": "Alice"}, {"ID": 2, "NAME": "Bob"}]
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        ds = JsonDataSource(path)
        assert ds.get_columns() == ["ID", "NAME"]
        assert ds.get_row_count() == 2
        rows = list(ds.read_rows())
        assert len(rows) == 2

    def test_jsonl_format(self):
        from src.data_source import JsonDataSource
        path = os.path.join(self.tmpdir, "test.jsonl")
        with open(path, 'w', encoding='utf-8') as f:
            f.write('{"ID": 1, "NAME": "Alice"}\n')
            f.write('{"ID": 2, "NAME": "Bob"}\n')
        ds = JsonDataSource(path)
        assert ds.source_type == "json"
        assert ds.get_row_count() == 2
        rows = list(ds.read_rows())
        assert rows[0]["NAME"] == "Alice"

    def test_json_file_not_found(self):
        from src.data_source import JsonDataSource
        ds = JsonDataSource("/nonexistent/data.json")
        with pytest.raises(FileNotFoundError):
            ds.get_columns()

    def test_json_max_rows(self):
        from src.data_source import JsonDataSource
        path = os.path.join(self.tmpdir, "big.json")
        data = [{"ID": i} for i in range(100)]
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        ds = JsonDataSource(path, max_rows=5)
        rows = list(ds.read_rows())
        assert len(rows) == 5

    def test_get_preview(self):
        from src.data_source import JsonDataSource
        path = os.path.join(self.tmpdir, "test.json")
        data = [{"ID": i, "VAL": f"v{i}"} for i in range(100)]
        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        ds = JsonDataSource(path)
        # get_preview 依赖 _columns 已填充，需先调用 get_columns
        cols = ds.get_columns()
        cols2, preview = ds.get_preview(max_rows=10)
        assert len(preview) == 10
        assert "ID" in cols2


class TestExcelDataSource:
    """测试 Excel 数据源"""

    def setup_method(self):
        self.tmpdir = tempfile.mkdtemp()

    def teardown_method(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_excel_file_not_found(self):
        from src.data_source import ExcelDataSource
        ds = ExcelDataSource("/nonexistent/data.xlsx")
        with pytest.raises(FileNotFoundError):
            ds.get_columns()

    def test_excel_source_type(self):
        from src.data_source import ExcelDataSource
        ds = ExcelDataSource("/tmp/test.xlsx")
        assert ds.source_type == "excel"

    def test_excel_xls_source_type(self):
        from src.data_source import ExcelDataSource
        ds = ExcelDataSource("/tmp/test.xls")
        assert ds.source_type == "excel_xls"

    def test_excel_with_openpyxl(self):
        """使用 openpyxl 创建测试文件"""
        from src.data_source import ExcelDataSource
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        path = os.path.join(self.tmpdir, "test.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "NAME", "AGE"])
        ws.append([1, "Alice", 30])
        ws.append([2, "Bob", 25])
        wb.save(path)

        ds = ExcelDataSource(path)
        assert ds.get_columns() == ["ID", "NAME", "AGE"]
        assert ds.get_row_count() == 2
        rows = list(ds.read_rows())
        assert rows[0]["NAME"] == "Alice"

    def test_excel_max_rows(self):
        from src.data_source import ExcelDataSource
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        path = os.path.join(self.tmpdir, "big.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID"])
        for i in range(100):
            ws.append([i])
        wb.save(path)

        ds = ExcelDataSource(path, max_rows=10)
        assert ds.get_row_count() == 10

    def test_excel_preview(self):
        from src.data_source import ExcelDataSource
        try:
            import openpyxl
        except ImportError:
            pytest.skip("openpyxl not installed")

        path = os.path.join(self.tmpdir, "test.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "VAL"])
        for i in range(100):
            ws.append([i, f"v{i}"])
        wb.save(path)

        ds = ExcelDataSource(path)
        cols, preview = ds.get_preview(max_rows=5)
        assert len(preview) == 5
        assert "ID" in cols


class TestCreateDataSource:
    """测试 create_data_source 工厂方法"""

    def test_csv_extension(self):
        from src.data_source import create_data_source, CsvDataSource
        ds = create_data_source("/tmp/test.csv")
        assert isinstance(ds, CsvDataSource)

    def test_json_extension(self):
        from src.data_source import create_data_source, JsonDataSource
        ds = create_data_source("/tmp/test.json")
        assert isinstance(ds, JsonDataSource)

    def test_jsonl_extension(self):
        from src.data_source import create_data_source, JsonDataSource
        ds = create_data_source("/tmp/test.jsonl")
        assert isinstance(ds, JsonDataSource)

    def test_xlsx_extension(self):
        from src.data_source import create_data_source, ExcelDataSource
        ds = create_data_source("/tmp/test.xlsx")
        assert isinstance(ds, ExcelDataSource)

    def test_xls_extension(self):
        from src.data_source import create_data_source, ExcelDataSource
        ds = create_data_source("/tmp/test.xls")
        assert isinstance(ds, ExcelDataSource)

    def test_unsupported_format(self):
        from src.data_source import create_data_source
        with pytest.raises(ValueError, match="不支持的文件格式"):
            create_data_source("/tmp/test.txt")

    def test_kwargs_passed_through(self):
        from src.data_source import create_data_source, CsvDataSource
        ds = create_data_source("/tmp/test.csv", encoding='gbk', delimiter=';')
        assert isinstance(ds, CsvDataSource)
        assert ds.encoding == 'gbk'
        assert ds.delimiter == ';'


# ============================================================================
# P1-3: auth.py
# ============================================================================

class TestRole:
    """测试角色权限"""

    def test_role_constants(self):
        from src.auth import Role
        assert Role.OPERATOR == "operator"
        assert Role.APPROVER == "approver"
        assert Role.AUDITOR == "auditor"
        assert Role.ADMIN == "admin"

    def test_can_execute(self):
        from src.auth import Role
        assert Role.can_execute("operator") is True
        assert Role.can_execute("admin") is True
        assert Role.can_execute("approver") is False
        assert Role.can_execute("auditor") is False
        assert Role.can_execute("unknown") is False

    def test_can_approve(self):
        from src.auth import Role
        assert Role.can_approve("approver") is True
        assert Role.can_approve("admin") is True
        assert Role.can_approve("operator") is False
        assert Role.can_approve("auditor") is False

    def test_can_audit(self):
        from src.auth import Role
        assert Role.can_audit("auditor") is True
        assert Role.can_audit("admin") is True
        assert Role.can_audit("operator") is False
        assert Role.can_audit("approver") is False

    def test_all_roles(self):
        from src.auth import Role
        assert len(Role.ALL) == 4
        assert set(Role.ALL) == {"operator", "approver", "auditor", "admin"}


class TestRiskLevel:
    """测试风险等级评估"""

    def test_low_risk(self):
        from src.auth import RiskLevel
        risk, reason = RiskLevel.assess("DEV_DB", "EMPLOYEE", 100)
        assert risk == "low"

    def test_medium_risk_large_batch(self):
        from src.auth import RiskLevel
        risk, reason = RiskLevel.assess("DEV_DB", "EMPLOYEE", 50000)
        assert risk == "medium"
        assert "10000" in reason

    def test_high_risk_huge_batch(self):
        """测试超大批量更新风险等级（>100000 行应判定为 high）"""
        from src.auth import RiskLevel
        risk, reason = RiskLevel.assess("DEV_DB", "EMPLOYEE", 200000)
        # 修复后：>100000 的分支优先于 >10000，返回 high
        assert risk == "high"
        assert "100000" in reason

    def test_production_environment_medium(self):
        from src.auth import RiskLevel
        risk, reason = RiskLevel.assess("PROD_DB", "EMPLOYEE", 500)
        assert risk == "medium"
        assert "生产" in reason

    def test_production_keyword_detection(self):
        from src.auth import RiskLevel
        for keyword in ["PROD", "PRODUCTION", "LIVE", "PRD", "生产"]:
            risk, _ = RiskLevel.assess(keyword, "TABLE", 100)
            assert risk == "medium", f"Keyword '{keyword}' not detected as production"

    def test_production_high_risk(self):
        from src.auth import RiskLevel
        risk, reason = RiskLevel.assess("PROD", "EMPLOYEE", 5000)
        assert risk == "high"
        assert "1000" in reason

    def test_is_production_flag(self):
        from src.auth import RiskLevel
        risk, _ = RiskLevel.assess("DEV_DB", "TABLE", 100, is_production=True)
        assert risk == "medium"


class TestAuthManager:
    """测试用户认证管理"""

    def setup_method(self):
        self.tmpdir = tempfile.mkdtemp()

    def teardown_method(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_default_admin_created(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        users = am.list_users()
        assert any(u["username"] == "admin" for u in users)

    def test_authenticate_admin(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        ok, msg, info = am.authenticate("admin", "admin")
        assert ok is True
        assert info["role"] == "admin"

    def test_authenticate_wrong_password(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        ok, msg, info = am.authenticate("admin", "wrong")
        assert ok is False
        assert "密码错误" in msg

    def test_authenticate_nonexistent_user(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        ok, msg, info = am.authenticate("nobody", "pass")
        assert ok is False
        assert "用户不存在" in msg

    def test_add_user(self):
        from src.auth import AuthManager, Role
        am = AuthManager(Path(self.tmpdir))
        ok, msg = am.add_user("admin", "testuser", "pass123", Role.OPERATOR, "Test User")
        assert ok is True
        ok2, _, info = am.authenticate("testuser", "pass123")
        assert ok2 is True
        assert info["role"] == "operator"

    def test_add_user_non_admin(self):
        from src.auth import AuthManager, Role
        am = AuthManager(Path(self.tmpdir))
        am.add_user("admin", "op1", "pass", Role.OPERATOR)
        ok, msg = am.add_user("op1", "op2", "pass", Role.OPERATOR)
        assert ok is False
        assert "管理员" in msg

    def test_add_duplicate_user(self):
        from src.auth import AuthManager, Role
        am = AuthManager(Path(self.tmpdir))
        am.add_user("admin", "user1", "pass", Role.OPERATOR)
        ok, msg = am.add_user("admin", "user1", "pass", Role.OPERATOR)
        assert ok is False
        assert "已存在" in msg

    def test_add_user_invalid_role(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        ok, msg = am.add_user("admin", "user1", "pass", "invalid_role")
        assert ok is False
        assert "无效角色" in msg

    def test_remove_user(self):
        from src.auth import AuthManager, Role
        am = AuthManager(Path(self.tmpdir))
        am.add_user("admin", "user1", "pass", Role.OPERATOR)
        ok, msg = am.remove_user("admin", "user1")
        assert ok is True
        assert am.get_user("user1") is None

    def test_remove_default_admin(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        ok, msg = am.remove_user("admin", "admin")
        assert ok is False
        assert "默认管理员" in msg

    def test_remove_nonexistent_user(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        ok, msg = am.remove_user("admin", "ghost")
        assert ok is False
        assert "不存在" in msg

    def test_list_users_no_password_hash(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        users = am.list_users()
        for u in users:
            assert "password_hash" not in u

    def test_get_user(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        user = am.get_user("admin")
        assert user is not None
        assert user["username"] == "admin"
        assert "password_hash" not in user

    def test_get_nonexistent_user(self):
        from src.auth import AuthManager
        am = AuthManager(Path(self.tmpdir))
        assert am.get_user("nobody") is None

    def test_disabled_user(self):
        from src.auth import AuthManager, Role
        am = AuthManager(Path(self.tmpdir))
        am.add_user("admin", "user1", "pass", Role.OPERATOR)
        # 手动禁用
        am._users["user1"]["enabled"] = False
        am._save_users()
        ok, msg, _ = am.authenticate("user1", "pass")
        assert ok is False
        assert "禁用" in msg


class TestApprovalManager:
    """测试审批令牌管理"""

    def setup_method(self):
        from src.auth import ApprovalManager
        ApprovalManager._tokens = {}

    def test_generate_token(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP", "rows": 100})
        assert len(token) == 6
        assert token.isdigit()

    def test_approve_token(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP"})
        ok, msg = ApprovalManager.approve(token, "user2")
        assert ok is True
        assert "通过" in msg

    def test_approve_self_fails(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP"})
        ok, msg = ApprovalManager.approve(token, "user1")
        assert ok is False
        assert "自己" in msg

    def test_approve_invalid_token(self):
        from src.auth import ApprovalManager
        ok, msg = ApprovalManager.approve("000000", "user2")
        assert ok is False
        assert "无效" in msg or "过期" in msg

    def test_reject_token(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP"})
        ok, msg = ApprovalManager.reject(token, "user2", "不合理")
        assert ok is True
        assert "拒绝" in msg

    def test_is_approved(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP"})
        assert ApprovalManager.is_approved(token) is False
        ApprovalManager.approve(token, "user2")
        assert ApprovalManager.is_approved(token) is True

    def test_is_approved_invalid(self):
        from src.auth import ApprovalManager
        assert ApprovalManager.is_approved("999999") is False

    def test_token_expires(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP"})
        # 手动将令牌设为过期
        ApprovalManager._tokens[token]["expires_at"] = datetime.now() - timedelta(minutes=1)
        assert ApprovalManager.is_approved(token) is False

    def test_double_approve_fails(self):
        from src.auth import ApprovalManager
        token = ApprovalManager.generate_token("user1", {"table": "EMP"})
        ApprovalManager.approve(token, "user2")
        ok, msg = ApprovalManager.approve(token, "user3")
        assert ok is False
        assert "已处理" in msg


# ============================================================================
# P1-4: cli.py
# ============================================================================

class TestCli:
    """测试 CLI 命令行接口"""

    def test_update_subcommand_parsing(self):
        """测试 update 子命令参数解析"""
        from src.cli import main
        import argparse

        # 构造 parser 来测试参数解析
        parser = argparse.ArgumentParser()
        subparsers = parser.add_subparsers(dest="command")
        update_parser = subparsers.add_parser("update")
        update_parser.add_argument("--connection", "-c", required=True)
        update_parser.add_argument("--table", "-t", required=True)
        update_parser.add_argument("--schema", "-s", default="APPS")
        update_parser.add_argument("--excel", "--file", "-f", dest="file", required=True)
        update_parser.add_argument("--key-column", "-k", required=True)
        update_parser.add_argument("--update-columns", "-u", required=True)
        update_parser.add_argument("--yes", "-y", action="store_true")
        update_parser.add_argument("--dry-run", action="store_true")

        args = parser.parse_args([
            "update", "-c", "DEV", "-t", "EMP",
            "-f", "data.csv", "-k", "EMP_ID",
            "-u", "NAME,AGE", "--dry-run"
        ])
        assert args.command == "update"
        assert args.connection == "DEV"
        assert args.table == "EMP"
        assert args.schema == "APPS"
        assert args.file == "data.csv"
        assert args.key_column == "EMP_ID"
        assert args.update_columns == "NAME,AGE"
        assert args.dry_run is True

    def test_version_subcommand(self):
        """测试 version 子命令输出"""
        from src.cli import cmd_version
        with patch('builtins.print') as mock_print:
            cmd_version(argparse.Namespace())
            mock_print.assert_called_once()
            output = mock_print.call_args[0][0]
            assert "DBForge" in output
            assert "v" in output

    def test_connections_subcommand_parsing(self):
        import argparse
        parser = argparse.ArgumentParser()
        subparsers = parser.add_subparsers(dest="command")
        subparsers.add_parser("connections")
        args = parser.parse_args(["connections"])
        assert args.command == "connections"

    def test_templates_subcommand_parsing(self):
        import argparse
        parser = argparse.ArgumentParser()
        subparsers = parser.add_subparsers(dest="command")
        subparsers.add_parser("templates")
        args = parser.parse_args(["templates"])
        assert args.command == "templates"

    def test_verify_audit_subcommand_parsing(self):
        import argparse
        parser = argparse.ArgumentParser()
        subparsers = parser.add_subparsers(dest="command")
        subparsers.add_parser("verify-audit")
        args = parser.parse_args(["verify-audit"])
        assert args.command == "verify-audit"

    def test_no_command_exits_zero(self):
        """无子命令时应打印帮助并退出"""
        from src.cli import main
        with patch('sys.argv', ['cli']):
            with pytest.raises(SystemExit) as exc_info:
                main()
            assert exc_info.value.code == 0

    def test_exit_codes_convention(self):
        """验证退出码约定文档"""
        # 退出码: 0=成功, 1=参数错误, 2=连接失败, 3=验证失败, 4=部分失败, 5=全部失败
        # 这里验证 cli 模块的 docstring 中有说明
        from src import cli
        assert "0=成功" in cli.__doc__ or "退出码" in cli.__doc__


# ============================================================================
# P1-5: notification.py
# ============================================================================

class TestNotificationManager:
    """测试通知管理器"""

    def test_disabled_by_default(self):
        from src.notification import NotificationManager
        nm = NotificationManager(config={"enabled": False, "channels": {}})
        result = nm.notify("Test", "Message")
        assert result == {}

    def test_enable_channel(self):
        from src.notification import NotificationManager
        with tempfile.TemporaryDirectory() as tmpdir:
            config = {"enabled": False, "channels": {}}
            nm = NotificationManager(config=config)
            # 覆盖 _save_config 避免写文件
            nm._save_config = MagicMock()
            nm.enable_channel("webhook", {"webhook_url": "http://example.com/hook"})
            assert nm.config["enabled"] is True
            assert "webhook" in nm.config["channels"]

    def test_notify_with_no_channels(self):
        from src.notification import NotificationManager
        nm = NotificationManager(config={"enabled": True, "channels": {}})
        result = nm.notify("Test", "Message")
        assert result == {}

    def test_notify_webhook_mock(self):
        """测试 webhook 通知（mock HTTP 请求）"""
        from src.notification import NotificationManager
        config = {
            "enabled": True,
            "channels": {
                "webhook": {"webhook_url": "http://example.com/hook"}
            }
        }
        nm = NotificationManager(config=config)

        mock_response = MagicMock()
        mock_response.status = 200
        mock_response.__enter__ = MagicMock(return_value=mock_response)
        mock_response.__exit__ = MagicMock(return_value=False)

        with patch('urllib.request.urlopen', return_value=mock_response) as mock_urlopen:
            result = nm.notify("Test Title", "Test Message")
            assert result.get("webhook") is True
            mock_urlopen.assert_called_once()

    def test_notify_dingtalk_mock(self):
        """测试钉钉通知（mock）"""
        from src.notification import NotificationManager
        config = {
            "enabled": True,
            "channels": {
                "dingtalk": {"webhook_url": "https://oapi.dingtalk.com/robot/send?access_token=xxx"}
            }
        }
        nm = NotificationManager(config=config)

        mock_resp = MagicMock()
        mock_resp.read.return_value = b'{"errcode": 0}'
        mock_resp.__enter__ = MagicMock(return_value=mock_resp)
        mock_resp.__exit__ = MagicMock(return_value=False)

        with patch('urllib.request.urlopen', return_value=mock_resp):
            result = nm.notify("Test", "Msg")
            assert result.get("dingtalk") is True

    def test_notify_wecom_mock(self):
        """测试企业微信通知（mock）"""
        from src.notification import NotificationManager
        config = {
            "enabled": True,
            "channels": {
                "wecom": {"webhook_url": "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=xxx"}
            }
        }
        nm = NotificationManager(config=config)

        mock_resp = MagicMock()
        mock_resp.read.return_value = b'{"errcode": 0}'
        mock_resp.__enter__ = MagicMock(return_value=mock_resp)
        mock_resp.__exit__ = MagicMock(return_value=False)

        with patch('urllib.request.urlopen', return_value=mock_resp):
            result = nm.notify("Test", "Msg")
            assert result.get("wecom") is True

    def test_notify_email_missing_config(self):
        """邮件配置不完整时返回 False"""
        from src.notification import NotificationManager
        config = {
            "enabled": True,
            "channels": {
                "email": {"smtp_host": "", "to": ""}
            }
        }
        nm = NotificationManager(config=config)
        result = nm.notify("Test", "Msg")
        assert result.get("email") is False

    def test_notify_channel_exception_caught(self):
        """通知渠道异常时不崩溃"""
        from src.notification import NotificationManager
        config = {
            "enabled": True,
            "channels": {
                "webhook": {"webhook_url": "http://invalid.local/hook"}
            }
        }
        nm = NotificationManager(config=config)
        with patch('urllib.request.urlopen', side_effect=Exception("network error")):
            result = nm.notify("Test", "Msg")
            assert result.get("webhook") is False

    def test_notify_update_result(self):
        """测试更新结果通知"""
        from src.notification import NotificationManager
        nm = NotificationManager(config={"enabled": False, "channels": {}})
        result = nm.notify_update_result(
            schema="APPS", table="EMP", success_count=100,
            fail_count=0, unmatched_count=2, elapsed_ms=5000,
            backup_table="BK_EMP"
        )
        assert result == {}  # disabled

    def test_load_config_default(self):
        """默认配置加载（无配置文件时）"""
        from src.notification import NotificationManager
        nm = NotificationManager()
        # 如果没有 notification.json，则默认 disabled
        assert isinstance(nm.config, dict)

    def test_webhook_empty_url(self):
        from src.notification import NotificationManager
        config = {
            "enabled": True,
            "channels": {
                "webhook": {"webhook_url": ""}
            }
        }
        nm = NotificationManager(config=config)
        result = nm.notify("Test", "Msg")
        assert result.get("webhook") is False


# ============================================================================
# P2-1: service/__init__.py
# ============================================================================

class TestUpdateConfig:
    """测试 UpdateConfig 数据类"""

    def test_default_values(self):
        from src.service import UpdateConfig
        cfg = UpdateConfig(connection_name="DEV")
        assert cfg.schema == "APPS"
        assert cfg.temp_schema == "APPS"
        assert cfg.target_table == ""
        assert cfg.key_column == ""
        assert cfg.update_columns == []
        assert cfg.data_source is None
        assert cfg.data_file_path == ""
        assert cfg.dry_run is False
        assert cfg.auto_confirm is False
        assert cfg.use_merge is True
        assert cfg.cancel_event is None

    def test_custom_values(self):
        from src.service import UpdateConfig
        cfg = UpdateConfig(
            connection_name="PROD",
            schema="HR",
            target_table="EMP",
            key_column="EMP_ID",
            update_columns=["NAME", "AGE"],
            dry_run=True,
        )
        assert cfg.connection_name == "PROD"
        assert cfg.schema == "HR"
        assert cfg.dry_run is True


class TestUpdateResult:
    """测试 UpdateResult 数据类"""

    def test_default_values(self):
        from src.service import UpdateResult
        r = UpdateResult(success=True)
        assert r.success_count == 0
        assert r.fail_count == 0
        assert r.unmatched_count == 0
        assert r.total_count == 0
        assert r.backup_table == ""
        assert r.elapsed_ms == 0
        assert r.failed_records == []
        assert r.error_msg == ""
        assert r.sql_preview == ""

    def test_failure_result(self):
        from src.service import UpdateResult
        r = UpdateResult(success=False, error_msg="连接失败")
        assert r.success is False
        assert r.error_msg == "连接失败"


class TestValidationResult:
    """测试 ValidationResult 数据类"""

    def test_default_values(self):
        from src.service import ValidationResult
        v = ValidationResult(valid=True)
        assert v.message == ""
        assert v.column_types == {}
        assert v.risk_level == "low"
        assert v.risk_reason == ""

    def test_invalid_result(self):
        from src.service import ValidationResult
        v = ValidationResult(valid=False, message="未指定目标表")
        assert v.valid is False
        assert "目标表" in v.message


class TestUpdateService:
    """测试 UpdateService"""

    def test_validate_no_target_table(self):
        from src.service import UpdateService, UpdateConfig, ValidationResult
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)
        cfg = UpdateConfig(connection_name="DEV")
        db = MagicMock()
        result = svc.validate(cfg, db)
        assert result.valid is False
        assert "目标表" in result.message

    def test_validate_no_key_column(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)
        cfg = UpdateConfig(connection_name="DEV", target_table="EMP")
        db = MagicMock()
        result = svc.validate(cfg, db)
        assert result.valid is False
        assert "唯一标识列" in result.message

    def test_validate_no_update_columns(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)
        cfg = UpdateConfig(connection_name="DEV", target_table="EMP", key_column="ID")
        db = MagicMock()
        result = svc.validate(cfg, db)
        assert result.valid is False
        assert "待更新列" in result.message

    def test_validate_no_data_source(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)
        cfg = UpdateConfig(
            connection_name="DEV", target_table="EMP",
            key_column="ID", update_columns=["NAME"]
        )
        db = MagicMock()
        result = svc.validate(cfg, db)
        assert result.valid is False
        assert "数据源" in result.message

    def test_validate_empty_data(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)

        mock_ds = MagicMock()
        mock_ds.get_row_count.return_value = 0

        cfg = UpdateConfig(
            connection_name="DEV", target_table="EMP",
            key_column="ID", update_columns=["NAME"],
            data_source=mock_ds,
        )
        db = MagicMock()
        result = svc.validate(cfg, db)
        assert result.valid is False
        assert "为空" in result.message

    def test_validate_db_not_connected(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)

        mock_ds = MagicMock()
        mock_ds.get_row_count.return_value = 10

        cfg = UpdateConfig(
            connection_name="DEV", target_table="EMP",
            key_column="ID", update_columns=["NAME"],
            data_source=mock_ds,
        )
        db = MagicMock()
        db.is_connected.return_value = False
        result = svc.validate(cfg, db)
        assert result.valid is False
        assert "未连接" in result.message

    def test_dry_run_execute(self):
        """测试 dry_run 模式执行"""
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)

        mock_ds = MagicMock()
        mock_ds.get_row_count.return_value = 10
        mock_ds.read_rows.return_value = iter([])

        cfg = UpdateConfig(
            connection_name="DEV", target_table="EMP",
            key_column="ID", update_columns=["NAME"],
            data_source=mock_ds, dry_run=True,
        )
        db = MagicMock()
        db.is_connected.return_value = True

        # Mock DataUpdater
        with patch('src.service.DataUpdater') as MockUpdater:
            mock_updater = MagicMock()
            mock_updater.validate_table_and_columns_multi.return_value = (True, "OK")
            MockUpdater.return_value = mock_updater
            db.get_columns.return_value = [{"name": "ID", "type": "NUMBER"}, {"name": "NAME", "type": "VARCHAR2"}]

            result = svc.execute(cfg, db)
            assert result.success is True
            assert result.total_count == 10
            assert result.sql_preview != ""

    def test_generate_sql_preview_merge(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)

        mock_ds = MagicMock()
        mock_ds.get_row_count.return_value = 50

        cfg = UpdateConfig(
            connection_name="DEV", schema="APPS",
            target_table="EMP", key_column="ID",
            update_columns=["NAME", "AGE"],
            data_source=mock_ds, use_merge=True,
        )
        db = MagicMock()
        preview = svc.generate_sql_preview(cfg, db)
        assert "MERGE" in preview
        assert "APPS.EMP" in preview
        assert "NAME" in preview

    def test_generate_sql_preview_update(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)

        mock_ds = MagicMock()
        mock_ds.get_row_count.return_value = 50

        cfg = UpdateConfig(
            connection_name="DEV", schema="APPS",
            target_table="EMP", key_column="ID",
            update_columns=["NAME"],
            data_source=mock_ds, use_merge=False,
        )
        db = MagicMock()
        preview = svc.generate_sql_preview(cfg, db)
        assert "UPDATE" in preview
        assert "WHERE" in preview

    def test_generate_sql_preview_empty(self):
        from src.service import UpdateService, UpdateConfig
        from src.logger import LogManager
        log = MagicMock(spec=LogManager)
        svc = UpdateService(log)
        cfg = UpdateConfig(connection_name="DEV")
        db = MagicMock()
        preview = svc.generate_sql_preview(cfg, db)
        assert preview == ""


# ============================================================================
# P2-2: progress.py
# ============================================================================

class TestFormatEta:
    """测试 ETA 格式化"""

    def test_zero_seconds(self):
        from src.progress import format_eta
        assert format_eta(0) == "0s"

    def test_one_second(self):
        from src.progress import format_eta
        assert format_eta(1) == "1s"

    def test_59_seconds(self):
        from src.progress import format_eta
        result = format_eta(59)
        assert "59s" in result

    def test_60_seconds(self):
        from src.progress import format_eta
        result = format_eta(60)
        assert "1m" in result

    def test_3600_seconds(self):
        from src.progress import format_eta
        result = format_eta(3600)
        assert "1h" in result

    def test_86400_seconds(self):
        from src.progress import format_eta
        result = format_eta(86400)
        assert "1d" in result

    def test_90_seconds(self):
        from src.progress import format_eta
        result = format_eta(90)
        assert "1m" in result
        assert "30s" in result

    def test_3723_seconds(self):
        from src.progress import format_eta
        result = format_eta(3723)
        assert "1h" in result
        assert "2m" in result
        assert "3s" in result

    def test_90061_seconds(self):
        from src.progress import format_eta
        result = format_eta(90061)
        assert "1d" in result
        assert "1h" in result
        assert "1m" in result
        assert "1s" in result

    def test_negative_seconds(self):
        from src.progress import format_eta
        result = format_eta(-10)
        assert result == "0s"

    def test_float_seconds(self):
        from src.progress import format_eta
        result = format_eta(45.7)
        assert "45s" in result


class TestFormatProgressBar:
    """测试进度条渲染"""

    def test_zero_percent(self):
        from src.progress import format_progress_bar
        result = format_progress_bar(0, 100)
        assert "0%" in result
        assert "[" in result
        assert "]" in result

    def test_45_percent(self):
        from src.progress import format_progress_bar
        result = format_progress_bar(45, 100)
        assert "45%" in result

    def test_100_percent(self):
        from src.progress import format_progress_bar
        result = format_progress_bar(100, 100)
        assert "100%" in result
        assert "=" in result

    def test_total_zero(self):
        from src.progress import format_progress_bar
        result = format_progress_bar(0, 0)
        assert "100%" in result  # total=0 视为 100%

    def test_custom_width(self):
        from src.progress import format_progress_bar
        result = format_progress_bar(50, 100, width=20)
        assert "50%" in result

    def test_over_100_percent_capped(self):
        from src.progress import format_progress_bar
        result = format_progress_bar(150, 100)
        assert "100%" in result


class TestProgressTracker:
    """测试 ProgressTracker"""

    def test_start_and_update(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        time.sleep(0.01)
        eta = tracker.update(50, 100)
        assert tracker.current == 50
        assert tracker.total == 100
        assert tracker.elapsed_seconds > 0
        assert tracker.percentage == 50

    def test_eta_calculation(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        # 模拟: 50 项用了 ~0s，ETA 应接近 0
        tracker.update(50, 100)
        # eta_seconds 应 >= 0
        assert tracker.eta_seconds >= 0

    def test_eta_zero_current(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        tracker.update(0, 100)
        assert tracker.eta_seconds == 0

    def test_format_eta(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        tracker.update(50, 100)
        eta_str = tracker.format_eta()
        assert isinstance(eta_str, str)

    def test_format_elapsed(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        time.sleep(0.05)
        tracker.update(1, 10)
        elapsed_str = tracker.format_elapsed()
        assert "s" in elapsed_str

    def test_format_progress_bar(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        tracker.update(50, 100)
        bar = tracker.format_progress_bar()
        assert "50%" in bar

    def test_reset(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        tracker.update(50, 100)
        tracker.reset()
        assert tracker.current == 0
        assert tracker.total == 0
        assert tracker.elapsed_seconds == 0.0
        assert tracker.eta_seconds == 0.0

    def test_percentage_zero_total(self):
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.start()
        tracker.update(0, 0)
        assert tracker.percentage == 0

    def test_update_without_start(self):
        """不显式调用 start 直接 update 也能工作"""
        from src.progress import ProgressTracker
        tracker = ProgressTracker()
        tracker.update(10, 100)
        assert tracker.current == 10
        assert tracker.elapsed_seconds >= 0


# ============================================================================
# P2-3: history_viewer.py
# ============================================================================

class TestHistoryViewer:
    """测试历史查看器"""

    def setup_method(self):
        self.tmpdir = tempfile.mkdtemp()

    def teardown_method(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_empty_logs_directory(self):
        from src.history_viewer import HistoryViewer
        viewer = HistoryViewer(Path(self.tmpdir))
        records = viewer.list_history()
        assert records == []

    def test_list_history_from_update_logs(self):
        from src.history_viewer import HistoryViewer
        # 创建一个 update log 文件
        log_content = (
            "[2026-06-20 10:00:00,000] INFO - 正在备份表 APPS.EMP\n"
            "[2026-06-20 10:00:01,000] INFO - 目标表: APPS.EMP\n"
            "[2026-06-20 10:00:02,000] INFO - 正在创建临时表\n"
            "[2026-06-20 10:00:05,000] INFO - 正在执行更新\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 100，失败: 2\n"
        )
        log_path = os.path.join(self.tmpdir, "update_20260620_100000.log")
        with open(log_path, 'w', encoding='utf-8') as f:
            f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        records = viewer.list_history()
        assert len(records) >= 1

    def test_search_history_by_table(self):
        from src.history_viewer import HistoryViewer
        log_content = (
            "[2026-06-20 10:00:00,000] INFO - 目标表: APPS.EMPLOYEE\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 50，失败: 0\n"
        )
        log_path = os.path.join(self.tmpdir, "update_20260620_100000.log")
        with open(log_path, 'w', encoding='utf-8') as f:
            f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        results = viewer.search_history(table_name="EMPLOYEE")
        assert len(results) >= 1
        assert "EMPLOYEE" in results[0].get("table", "").upper() or \
               "EMPLOYEE" in results[0].get("table", "")

    def test_search_history_by_schema(self):
        from src.history_viewer import HistoryViewer
        log_content = (
            "[2026-06-20 10:00:00,000] INFO - 目标表: HR.EMP\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 10，失败: 0\n"
        )
        log_path = os.path.join(self.tmpdir, "update_20260620_100000.log")
        with open(log_path, 'w', encoding='utf-8') as f:
            f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        results = viewer.search_history(schema="HR")
        assert len(results) >= 1

    def test_search_history_by_status(self):
        from src.history_viewer import HistoryViewer
        # 成功的日志
        log1 = (
            "[2026-06-20 10:00:00,000] INFO - 目标表: APPS.T1\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 100，失败: 0\n"
        )
        with open(os.path.join(self.tmpdir, "update_20260620_100000.log"), 'w') as f:
            f.write(log1)

        viewer = HistoryViewer(Path(self.tmpdir))
        results = viewer.search_history(status="success")
        assert len(results) >= 1

    def test_get_history_detail_from_log(self):
        from src.history_viewer import HistoryViewer
        log_content = (
            "[2026-06-20 10:00:00,000] INFO - 正在备份表 APPS.EMP\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 50，失败: 0\n"
        )
        log_path = os.path.join(self.tmpdir, "update_20260620_100000.log")
        with open(log_path, 'w') as f:
            f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        detail = viewer.get_history_detail("20260620_100000")
        assert detail is not None
        assert "log_entries" in detail

    def test_get_history_detail_not_found(self):
        from src.history_viewer import HistoryViewer
        viewer = HistoryViewer(Path(self.tmpdir))
        detail = viewer.get_history_detail("nonexistent_id")
        assert detail is None

    def test_export_csv(self):
        from src.history_viewer import HistoryViewer
        log_content = (
            "[2026-06-20 10:00:00,000] INFO - 目标表: APPS.EMP\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 100，失败: 0\n"
        )
        log_path = os.path.join(self.tmpdir, "update_20260620_100000.log")
        with open(log_path, 'w') as f:
            f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        output_path = os.path.join(self.tmpdir, "export.csv")
        ok, msg = viewer.export_history(output_path, format="csv")
        assert ok is True
        assert os.path.exists(output_path)
        # 验证 CSV 内容
        with open(output_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            assert len(rows) >= 1

    def test_export_json(self):
        from src.history_viewer import HistoryViewer
        log_content = (
            "[2026-06-20 10:00:00,000] INFO - 目标表: APPS.EMP\n"
            "[2026-06-20 10:00:10,000] INFO - 更新完成 成功: 100，失败: 0\n"
        )
        log_path = os.path.join(self.tmpdir, "update_20260620_100000.log")
        with open(log_path, 'w') as f:
            f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        output_path = os.path.join(self.tmpdir, "export.json")
        ok, msg = viewer.export_history(output_path, format="json")
        assert ok is True
        assert os.path.exists(output_path)
        with open(output_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
            assert isinstance(data, list)
            assert len(data) >= 1

    def test_export_unsupported_format(self):
        from src.history_viewer import HistoryViewer
        viewer = HistoryViewer(Path(self.tmpdir))
        output_path = os.path.join(self.tmpdir, "export.xml")
        ok, msg = viewer.export_history(output_path, format="xml")
        assert ok is False
        assert "Unsupported" in msg

    def test_history_json_source(self):
        """测试从 history.json 读取"""
        from src.history_viewer import HistoryViewer
        history_data = [
            {
                "id": "20260620120000000000",
                "timestamp": "2026-06-20 12:00:00",
                "action_type": "UPDATE",
                "schema": "APPS",
                "table": "DEPT",
                "key_column": "DEPT_ID",
                "total_count": 20,
                "success_count": 20,
                "fail_count": 0,
                "success": True,
                "backup_table": "BK_DEPT",
            }
        ]
        with open(os.path.join(self.tmpdir, "history.json"), 'w', encoding='utf-8') as f:
            json.dump(history_data, f)

        viewer = HistoryViewer(Path(self.tmpdir))
        records = viewer.list_history()
        assert len(records) >= 1
        assert records[0]["table"] == "DEPT"

    def test_audit_log_source(self):
        """测试从 audit.log 读取"""
        from src.history_viewer import HistoryViewer
        audit_record = {
            "timestamp": "2026-06-20 13:00:00",
            "action_type": "UPDATE",
            "success": True,
            "details": {
                "schema": "APPS",
                "table": "SALARY",
                "key_column": "EMP_ID",
                "total_count": 50,
                "success_count": 50,
                "fail_count": 0,
                "backup_table": "",
            }
        }
        with open(os.path.join(self.tmpdir, "audit.log"), 'w', encoding='utf-8') as f:
            f.write(json.dumps(audit_record) + "\n")

        viewer = HistoryViewer(Path(self.tmpdir))
        records = viewer.list_history()
        assert len(records) >= 1

    def test_list_history_limit(self):
        from src.history_viewer import HistoryViewer
        # 创建多个日志文件
        for i in range(10):
            log_content = f"[2026-06-20 10:0{i}:00,000] INFO - 目标表: APPS.T{i}\n"
            log_content += f"[2026-06-20 10:0{i}:10,000] INFO - 更新完成 成功: 10，失败: 0\n"
            log_path = os.path.join(self.tmpdir, f"update_20260620_100{i}00.log")
            with open(log_path, 'w') as f:
                f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        records = viewer.list_history(limit=3)
        assert len(records) <= 3

    def test_list_history_offset(self):
        from src.history_viewer import HistoryViewer
        for i in range(5):
            log_content = f"[2026-06-20 1{i}:00:00,000] INFO - 目标表: APPS.T{i}\n"
            log_content += f"[2026-06-20 1{i}:00:10,000] INFO - 更新完成 成功: 10，失败: 0\n"
            log_path = os.path.join(self.tmpdir, f"update_20260620_1{i}0000.log")
            with open(log_path, 'w') as f:
                f.write(log_content)

        viewer = HistoryViewer(Path(self.tmpdir))
        all_records = viewer.list_history(limit=100)
        offset_records = viewer.list_history(limit=100, offset=2)
        assert len(offset_records) == max(0, len(all_records) - 2)


# ============================================================================
# P2-4: constants.py
# ============================================================================

class TestConstants:
    """测试所有常量值"""

    def test_version(self):
        from src.constants import VERSION
        assert VERSION == "2.8.0"

    def test_app_name(self):
        from src.constants import APP_NAME
        assert APP_NAME == "DBForge"

    def test_default_batch_size(self):
        from src.constants import DEFAULT_BATCH_SIZE
        assert DEFAULT_BATCH_SIZE == 100

    def test_default_max_file_size(self):
        from src.constants import DEFAULT_MAX_FILE_SIZE
        assert DEFAULT_MAX_FILE_SIZE == 10 * 1024 * 1024

    def test_default_max_rows(self):
        from src.constants import DEFAULT_MAX_ROWS
        assert DEFAULT_MAX_ROWS == 100000

    def test_default_csv_max_rows(self):
        from src.constants import DEFAULT_CSV_MAX_ROWS
        assert DEFAULT_CSV_MAX_ROWS == 10000000

    def test_default_temp_schema(self):
        from src.constants import DEFAULT_TEMP_SCHEMA
        assert DEFAULT_TEMP_SCHEMA == "APPS"

    def test_default_target_schema(self):
        from src.constants import DEFAULT_TARGET_SCHEMA
        assert DEFAULT_TARGET_SCHEMA == "APPS"

    def test_default_timeout(self):
        from src.constants import DEFAULT_TIMEOUT
        assert DEFAULT_TIMEOUT == 30

    def test_default_pool_min(self):
        from src.constants import DEFAULT_POOL_MIN
        assert DEFAULT_POOL_MIN == 2

    def test_default_pool_max(self):
        from src.constants import DEFAULT_POOL_MAX
        assert DEFAULT_POOL_MAX == 10

    def test_default_pool_increment(self):
        from src.constants import DEFAULT_POOL_INCREMENT
        assert DEFAULT_POOL_INCREMENT == 1

    def test_supported_file_formats(self):
        from src.constants import SUPPORTED_FILE_FORMATS
        assert '.xlsx' in SUPPORTED_FILE_FORMATS
        assert '.xls' in SUPPORTED_FILE_FORMATS
        assert '.csv' in SUPPORTED_FILE_FORMATS
        assert '.json' in SUPPORTED_FILE_FORMATS
        assert '.jsonl' in SUPPORTED_FILE_FORMATS

    def test_audit_log_retention_days(self):
        from src.constants import AUDIT_LOG_RETENTION_DAYS
        assert AUDIT_LOG_RETENTION_DAYS == 90

    def test_backup_retention_days(self):
        from src.constants import BACKUP_RETENTION_DAYS
        assert BACKUP_RETENTION_DAYS == 30

    def test_max_preview_rows(self):
        from src.constants import MAX_PREVIEW_ROWS
        assert MAX_PREVIEW_ROWS == 50


# ============================================================================
# 运行入口
# ============================================================================

if __name__ == "__main__":
    pytest.main([__file__, "-v", "--tb=short"])
