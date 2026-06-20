# v2.9.1
#!/usr/bin/env python3
"""
核心模块深度功能测试
覆盖: db_connection.py, data_updater.py, config_manager.py, excel_handler.py
"""

import unittest
import os
import sys
import json
import tempfile
import shutil
import time
import threading
from unittest.mock import MagicMock, Mock, patch, PropertyMock
from datetime import datetime
from pathlib import Path

# ---------------------------------------------------------------------------
# 确保 src 可导入
# ---------------------------------------------------------------------------
sys.path.insert(0, str(Path(__file__).parent))

from src.db_connection import DBConnection, ConnectionPool
from src.data_updater import DataUpdater
from src.config_manager import ConfigManager, password_encrypt, password_decrypt
from src.excel_handler import ExcelHandler, MAX_FILE_SIZE, MAX_ROWS, MAX_PREVIEW_ROWS
from src.logger import LogManager
from src.errors import CancelledError
from src.security import SecurityError


# ============================================================================
# 1. db_connection.py — DBConnection 和 ConnectionPool
# ============================================================================

class TestDBConnectionConnect(unittest.TestCase):
    """测试 DBConnection.connect 方法签名与行为"""

    def setUp(self):
        self.db = DBConnection()

    def test_connect_signature(self):
        """connect 方法接受位置参数 (host, port, service, username, password) + 可选 db_type/database"""
        import inspect
        sig = inspect.signature(self.db.connect)
        params = list(sig.parameters.keys())
        # v2.9.0 增加 db_type/database 关键字参数（有默认值），保持向后兼容
        self.assertEqual(params[:5], ["host", "port", "service", "username", "password"])
        self.assertIn("db_type", params)
        self.assertIn("database", params)

    @patch("src.db_connection.oracledb")
    def test_connect_success(self, mock_oracledb):
        """连接成功时返回 (True, '连接成功')"""
        mock_oracledb.makedsn.return_value = "fake_dsn"
        mock_oracledb.connect.return_value = MagicMock()
        ok, msg = self.db.connect("localhost", 1521, "ORCL", "user1", "pass1")
        self.assertTrue(ok)
        self.assertEqual(msg, "连接成功")
        self.assertIsNotNone(self.db.connection)
        self.assertEqual(self.db.connection_info["host"], "localhost")
        self.assertEqual(self.db.connection_info["port"], 1521)

    @patch("src.db_connection.oracledb")
    def test_connect_ora_12541(self, mock_oracledb):
        """ORA-12541 错误返回 TNS 无监听提示"""
        import oracledb
        mock_oracledb.DatabaseError = oracledb.DatabaseError
        mock_oracledb.makedsn.return_value = "fake_dsn"
        mock_oracledb.connect.side_effect = oracledb.DatabaseError("ORA-12541: TNS:no listener")
        # 让 translate_ora_error 走 fallback 分支
        with patch("src.db_connection.translate_ora_error", side_effect=lambda e: f"数据库错误: {e}"):
            ok, msg = self.db.connect("badhost", 1521, "ORCL", "u", "p")
        self.assertFalse(ok)
        self.assertIn("TNS", msg)

    @patch("src.db_connection.oracledb")
    def test_connect_ora_01017(self, mock_oracledb):
        """ORA-01017 错误返回用户名/密码无效提示"""
        import oracledb
        mock_oracledb.DatabaseError = oracledb.DatabaseError
        mock_oracledb.makedsn.return_value = "fake_dsn"
        mock_oracledb.connect.side_effect = oracledb.DatabaseError("ORA-01017: invalid username/password")
        with patch("src.db_connection.translate_ora_error", side_effect=lambda e: f"数据库错误: {e}"):
            ok, msg = self.db.connect("localhost", 1521, "ORCL", "bad", "bad")
        self.assertFalse(ok)
        self.assertIn("用户名或密码", msg)

    @patch("src.db_connection.oracledb")
    def test_connect_generic_exception(self, mock_oracledb):
        """非 DatabaseError 异常也能被捕获"""
        import oracledb as real_oracledb
        # 保持 DatabaseError 为真实异常类，否则 except 子句会报 TypeError
        mock_oracledb.DatabaseError = real_oracledb.DatabaseError
        mock_oracledb.makedsn.side_effect = RuntimeError("unexpected")
        ok, msg = self.db.connect("localhost", 1521, "ORCL", "u", "p")
        self.assertFalse(ok)
        self.assertIn("unexpected", msg)


class TestDBConnectionDisconnect(unittest.TestCase):
    """测试 disconnect 方法"""

    def test_disconnect_when_connected(self):
        db = DBConnection()
        mock_conn = MagicMock()
        db.connection = mock_conn
        db.disconnect()
        mock_conn.close.assert_called_once()
        self.assertIsNone(db.connection)

    def test_disconnect_when_not_connected(self):
        db = DBConnection()
        db.connection = None
        # 不应抛异常
        db.disconnect()
        self.assertIsNone(db.connection)

    def test_disconnect_close_exception_swallowed(self):
        db = DBConnection()
        mock_conn = MagicMock()
        mock_conn.close.side_effect = Exception("close failed")
        db.connection = mock_conn
        # 异常被吞掉
        db.disconnect()
        self.assertIsNone(db.connection)


class TestDBConnectionExecuteSql(unittest.TestCase):
    """测试 execute_sql 方法参数验证"""

    def setUp(self):
        self.db = DBConnection()

    def test_execute_sql_not_connected(self):
        """未连接时返回错误"""
        ok, result, err = self.db.execute_sql("SELECT 1 FROM DUAL")
        self.assertFalse(ok)
        self.assertIsNone(result)
        self.assertIn("未连接数据库", err)

    @patch("src.db_connection.oracledb")
    def test_execute_sql_with_params(self, mock_oracledb):
        """带参数执行 SQL"""
        mock_conn = MagicMock()
        mock_cursor = MagicMock()
        mock_cursor.description = [("ID",), ("NAME",)]
        mock_cursor.fetchall.return_value = [(1, "test")]
        mock_conn.cursor.return_value = mock_cursor
        self.db.connection = mock_conn

        ok, result, err = self.db.execute_sql(
            "SELECT * FROM t WHERE id = :id", {"id": 1}
        )
        self.assertTrue(ok)
        self.assertEqual(result, [(1, "test")])
        mock_cursor.execute.assert_called_once_with(
            "SELECT * FROM t WHERE id = :id", {"id": 1}
        )

    @patch("src.db_connection.oracledb")
    def test_execute_sql_without_params(self, mock_oracledb):
        """不带参数执行 SQL"""
        mock_conn = MagicMock()
        mock_cursor = MagicMock()
        mock_cursor.description = None
        mock_cursor.fetchall.return_value = []
        mock_conn.cursor.return_value = mock_cursor
        self.db.connection = mock_conn

        ok, result, err = self.db.execute_sql("UPDATE t SET x=1")
        self.assertTrue(ok)
        self.assertIsNone(result)  # cursor.description is None -> result is None

    @patch("src.db_connection.oracledb")
    def test_execute_sql_commit_true(self, mock_oracledb):
        """commit=True 时调用 connection.commit()"""
        mock_conn = MagicMock()
        mock_cursor = MagicMock()
        mock_cursor.description = None
        mock_conn.cursor.return_value = mock_cursor
        self.db.connection = mock_conn

        self.db.execute_sql("UPDATE t SET x=1", commit=True)
        mock_conn.commit.assert_called_once()

    @patch("src.db_connection.oracledb")
    def test_execute_sql_commit_false(self, mock_oracledb):
        """commit=False 时不调用 connection.commit()"""
        mock_conn = MagicMock()
        mock_cursor = MagicMock()
        mock_cursor.description = None
        mock_conn.cursor.return_value = mock_cursor
        self.db.connection = mock_conn

        self.db.execute_sql("UPDATE t SET x=1", commit=False)
        mock_conn.commit.assert_not_called()


class TestDBConnectionCheckPermissions(unittest.TestCase):
    """测试 check_permissions 方法"""

    def setUp(self):
        self.db = DBConnection()

    def test_check_permissions_not_connected(self):
        """未连接时返回缺少 CONNECTION 权限"""
        ok, msg, missing = self.db.check_permissions("EMPLOYEE")
        self.assertFalse(ok)
        self.assertIn("未连接数据库", msg)
        self.assertIn("CONNECTION", missing)

    @patch("src.db_connection.sanitize_identifier", side_effect=lambda n, c="": n)
    def test_check_permissions_all_granted(self, mock_sanitize):
        """拥有所有权限时返回 True"""
        mock_conn = MagicMock()
        # cursor1: session_privs CREATE TABLE -> 1
        # cursor2: user_tables -> 1 (owns the table)
        cursors = []
        for _ in range(4):
            c = MagicMock()
            cursors.append(c)
        cursors[0].fetchone.return_value = (1,)  # CREATE TABLE in session_privs
        cursors[1].fetchone.return_value = (1,)  # table exists in user_tables
        mock_conn.cursor.side_effect = cursors
        self.db.connection = mock_conn

        ok, msg, missing = self.db.check_permissions("EMPLOYEE")
        self.assertTrue(ok)
        self.assertEqual(missing, [])

    @patch("src.db_connection.sanitize_identifier", side_effect=lambda n, c="": n)
    def test_check_permissions_missing_create_table(self, mock_sanitize):
        """缺少 CREATE TABLE 权限"""
        mock_conn = MagicMock()
        cursors = []
        for _ in range(4):
            c = MagicMock()
            cursors.append(c)
        cursors[0].fetchone.return_value = (0,)  # no CREATE TABLE in session_privs
        cursors[1].fetchone.return_value = (0,)  # no CREATE TABLE in user_sys_privs
        cursors[2].fetchone.return_value = (1,)  # owns table
        mock_conn.cursor.side_effect = cursors
        self.db.connection = mock_conn

        ok, msg, missing = self.db.check_permissions("EMPLOYEE")
        self.assertFalse(ok)
        self.assertIn("CREATE TABLE", missing)


class TestConnectionPool(unittest.TestCase):
    """测试 ConnectionPool 的 acquire/release/close"""

    @patch("src.db_connection.oracledb")
    def test_pool_creation_success(self, mock_oracledb):
        """连接池创建成功"""
        mock_oracledb.makedsn.return_value = "dsn"
        mock_pool = MagicMock()
        mock_oracledb.create_pool.return_value = mock_pool

        pool = ConnectionPool("host", 1521, "ORCL", "u", "p", min_size=2, max_size=5)
        self.assertTrue(pool.is_pool_enabled)
        mock_oracledb.create_pool.assert_called_once()

    @patch("src.db_connection.oracledb")
    def test_pool_acquire_release(self, mock_oracledb):
        """acquire 和 release 正常调用"""
        mock_oracledb.makedsn.return_value = "dsn"
        mock_pool = MagicMock()
        mock_conn = MagicMock()
        mock_pool.acquire.return_value = mock_conn
        mock_oracledb.create_pool.return_value = mock_pool

        pool = ConnectionPool("host", 1521, "ORCL", "u", "p")
        conn = pool.acquire()
        self.assertEqual(conn, mock_conn)
        mock_pool.acquire.assert_called_once()

        pool.release(conn)
        mock_pool.release.assert_called_once_with(conn)

    @patch("src.db_connection.oracledb")
    def test_pool_get_connection_alias(self, mock_oracledb):
        """get_connection 是 acquire 的别名"""
        mock_oracledb.makedsn.return_value = "dsn"
        mock_pool = MagicMock()
        mock_conn = MagicMock()
        mock_pool.acquire.return_value = mock_conn
        mock_oracledb.create_pool.return_value = mock_pool

        pool = ConnectionPool("host", 1521, "ORCL", "u", "p")
        conn = pool.get_connection()
        self.assertEqual(conn, mock_conn)

    @patch("src.db_connection.oracledb")
    def test_pool_close(self, mock_oracledb):
        """close 方法关闭连接池"""
        mock_oracledb.makedsn.return_value = "dsn"
        mock_pool = MagicMock()
        mock_oracledb.create_pool.return_value = mock_pool

        pool = ConnectionPool("host", 1521, "ORCL", "u", "p")
        pool.close()
        mock_pool.close.assert_called_once()
        self.assertFalse(pool.is_pool_enabled)

    @patch("src.db_connection.oracledb")
    def test_pool_creation_failure_degrades(self, mock_oracledb):
        """连接池创建失败时降级"""
        import oracledb
        mock_oracledb.makedsn.return_value = "dsn"
        mock_oracledb.DatabaseError = oracledb.DatabaseError
        mock_oracledb.create_pool.side_effect = oracledb.DatabaseError("pool failed")

        pool = ConnectionPool("host", 1521, "ORCL", "u", "p")
        self.assertFalse(pool.is_pool_enabled)

    @patch("src.db_connection.oracledb")
    def test_pool_acquire_failure_returns_none(self, mock_oracledb):
        """acquire 失败时返回 None"""
        import oracledb
        mock_oracledb.makedsn.return_value = "dsn"
        mock_pool = MagicMock()
        mock_pool.acquire.side_effect = oracledb.DatabaseError("acquire failed")
        mock_oracledb.DatabaseError = oracledb.DatabaseError
        mock_oracledb.create_pool.return_value = mock_pool

        pool = ConnectionPool("host", 1521, "ORCL", "u", "p")
        conn = pool.acquire()
        self.assertIsNone(conn)


# ============================================================================
# 2. data_updater.py — DataUpdater
# ============================================================================

class TestDataUpdaterCreateTempTable(unittest.TestCase):
    """测试 create_temp_table_multi_column 方法"""

    def setUp(self):
        self.mock_db = MagicMock()
        self.mock_log = MagicMock()
        self.updater = DataUpdater(self.mock_db, self.mock_log)

    def test_create_temp_table_success(self):
        """创建临时表成功"""
        self.mock_db.execute_sql.return_value = (True, None, "")
        self.mock_db.get_columns.return_value = []

        ok, temp_name = self.updater.create_temp_table_multi_column(
            "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )
        self.assertTrue(ok)
        self.assertTrue(temp_name.startswith("TEMP_UPDATE_"))
        self.assertTrue(self.updater.temp_table_created)

    def test_create_temp_table_with_type_inference(self):
        """P1-2: 从目标表推断列类型"""
        self.mock_db.execute_sql.return_value = (True, None, "")
        self.mock_db.get_columns.return_value = [
            {"name": "EMP_ID", "type": "NUMBER", "length": 22},
            {"name": "NAME", "type": "VARCHAR2", "length": 100},
            {"name": "AGE", "type": "NUMBER", "length": 22},
        ]

        ok, temp_name = self.updater.create_temp_table_multi_column(
            "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )
        self.assertTrue(ok)

        # 验证 CREATE TABLE SQL 使用了推断的类型
        create_sql = self.mock_db.execute_sql.call_args[0][0]
        self.assertIn("EMP_ID NUMBER", create_sql)
        self.assertIn("NAME VARCHAR2(100)", create_sql)
        self.assertIn("AGE NUMBER", create_sql)

    def test_create_temp_table_failure(self):
        """创建临时表失败"""
        self.mock_db.execute_sql.return_value = (False, None, "权限不足")
        self.mock_db.get_columns.return_value = []

        ok, err = self.updater.create_temp_table_multi_column(
            "APPS", "EMPLOYEE", "EMP_ID", ["NAME"]
        )
        self.assertFalse(ok)

    def test_create_temp_table_different_schema(self):
        """支持指定不同的临时表 Schema"""
        self.mock_db.execute_sql.return_value = (True, None, "")
        self.mock_db.get_columns.return_value = []

        ok, temp_name = self.updater.create_temp_table_multi_column(
            "SYSTEM", "EMPLOYEE", "EMP_ID", ["NAME"]
        )
        self.assertTrue(ok)
        create_sql = self.mock_db.execute_sql.call_args[0][0]
        self.assertIn("SYSTEM.TEMP_UPDATE_", create_sql)


class TestDataUpdaterExecuteMultiColumnUpdate(unittest.TestCase):
    """测试 execute_multi_column_update 方法"""

    def setUp(self):
        self.mock_db = MagicMock()
        self.mock_log = MagicMock()
        self.updater = DataUpdater(self.mock_db, self.mock_log)
        self.updater.temp_table_name = "TEMP_UPDATE_TEST"

    def test_successful_update(self):
        """正常更新流程"""
        # 行结构: key, old_NAME, old_AGE, cur_NAME, cur_AGE, new_NAME, new_AGE
        mock_result = [
            ("1001", "张三", 25, "张三", 25, "李四", 30)
        ]
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",)], None),    # 获取临时表 keys
            (True, mock_result, None),     # 查询匹配记录
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 1
        self.mock_db.connection.cursor.return_value = mock_cursor

        success, fail, records = self.updater.execute_multi_column_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )
        self.assertEqual(success, 1)
        self.assertEqual(fail, 0)

    def test_empty_fields_skipped(self):
        """空字段不更新"""
        # new_NAME="李四", new_AGE=None (空)
        mock_result = [
            ("1001", "张三", 25, "张三", 25, "李四", None)
        ]
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",)], None),
            (True, mock_result, None),
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 1
        self.mock_db.connection.cursor.return_value = mock_cursor

        success, fail, _ = self.updater.execute_multi_column_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )
        self.assertEqual(success, 1)
        # 验证 UPDATE SQL 只包含 NAME，不包含 AGE
        update_sql = mock_cursor.execute.call_args[0][0]
        self.assertIn("NAME = :NAME", update_sql)
        self.assertNotIn("AGE", update_sql)

    def test_all_fields_empty_counts_success(self):
        """所有字段为空时计为成功但不执行 UPDATE"""
        mock_result = [
            ("1001", "张三", 25, "张三", 25, None, None)
        ]
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",)], None),
            (True, mock_result, None),
        ]
        mock_cursor = MagicMock()
        self.mock_db.connection.cursor.return_value = mock_cursor

        success, fail, _ = self.updater.execute_multi_column_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )
        self.assertEqual(success, 1)
        self.assertEqual(fail, 0)
        # 没有执行 UPDATE
        mock_cursor.execute.assert_not_called()

    def test_unmatched_records(self):
        """未匹配记录被正确报告"""
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",), ("9999",)], None),  # 临时表 keys
            (True, [("1001", "A", "B", "A", "B", "C", "D")], None),  # 只有 1001 匹配
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 1
        self.mock_db.connection.cursor.return_value = mock_cursor

        success, fail, records = self.updater.execute_multi_column_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["COL1", "COL2"]
        )
        self.assertEqual(success, 1)
        # records 应包含 9999 的未匹配记录
        unmatched = [r for r in records if r["key_value"] == "9999"]
        self.assertEqual(len(unmatched), 1)
        self.assertIn("不存在", unmatched[0]["reason"])

    def test_no_matching_records(self):
        """没有匹配记录时返回空结果"""
        self.mock_db.execute_sql.side_effect = [
            (True, [("9998",), ("9999",)], None),
            (True, [], None),
        ]

        success, fail, records = self.updater.execute_multi_column_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME"]
        )
        self.assertEqual(success, 0)
        self.assertEqual(fail, 0)
        self.assertEqual(len(records), 2)  # 两条未匹配


class TestDataUpdaterExecuteMergeUpdate(unittest.TestCase):
    """测试 execute_merge_update 方法"""

    def setUp(self):
        self.mock_db = MagicMock()
        self.mock_log = MagicMock()
        self.updater = DataUpdater(self.mock_db, self.mock_log)
        self.updater.temp_table_name = "TEMP_UPDATE_TEST"

    def test_merge_update_success(self):
        """MERGE INTO 正常执行"""
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",), ("1002",)], None),  # 获取临时表 keys
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 2
        self.mock_db.connection.cursor.return_value = mock_cursor

        success, fail, records = self.updater.execute_merge_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )
        self.assertEqual(success, 2)
        self.assertEqual(fail, 0)

        # 验证执行了 MERGE INTO SQL
        merge_sql = mock_cursor.execute.call_args[0][0]
        self.assertIn("MERGE INTO", merge_sql)
        self.assertIn("APPS.EMPLOYEE", merge_sql)

    def test_merge_update_with_unmatched(self):
        """MERGE INTO 有未匹配记录"""
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",), ("9999",)], None),  # 临时表 keys
            (True, [("1001",), ("1002",)], None),  # 目标表 keys (用于检测未匹配)
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 1  # 只匹配了 1 条
        self.mock_db.connection.cursor.return_value = mock_cursor

        success, fail, records = self.updater.execute_merge_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME"]
        )
        self.assertEqual(success, 1)
        # 9999 应该被报告为未匹配
        unmatched = [r for r in records if r["key_value"] == "9999"]
        self.assertEqual(len(unmatched), 1)

    def test_merge_update_empty_columns_preserved(self):
        """MERGE INTO 空字段保留原值 (CASE WHEN)"""
        self.mock_db.execute_sql.side_effect = [
            (True, [("1001",)], None),
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 1
        self.mock_db.connection.cursor.return_value = mock_cursor

        self.updater.execute_merge_update(
            "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["NAME", "AGE"]
        )

        merge_sql = mock_cursor.execute.call_args[0][0]
        # 验证 SET 子句包含 CASE WHEN 逻辑
        self.assertIn("CASE WHEN", merge_sql)
        self.assertIn("IS NOT NULL", merge_sql)


class TestDataUpdaterCancel(unittest.TestCase):
    """测试 cancel 方法"""

    def setUp(self):
        self.mock_db = MagicMock()
        self.mock_log = MagicMock()
        self.updater = DataUpdater(self.mock_db, self.mock_log)

    def test_cancel_sets_event(self):
        """cancel 方法设置取消事件"""
        self.assertFalse(self.updater.cancel_event.is_set())
        self.updater.cancel()
        self.assertTrue(self.updater.cancel_event.is_set())
        self.mock_log.warning.assert_called()

    def test_cancel_during_update_raises_cancelled(self):
        """在更新过程中取消操作会抛出 CancelledError"""
        self.updater.temp_table_name = "TEMP_TEST"

        # 构造 100 条数据
        # 行结构: key(0) | old_COL1(1) | cur_COL1(2) | new_COL1(3)
        mock_result = [(str(i), "old", "cur", "new") for i in range(100)]
        self.mock_db.execute_sql.side_effect = [
            (True, [(str(i),) for i in range(100)], None),  # 临时表 keys
            (True, mock_result, None),                       # 匹配记录
        ]
        mock_cursor = MagicMock()
        mock_cursor.rowcount = 1
        self.mock_db.connection.cursor.return_value = mock_cursor

        # 在第一次 cursor.execute 时设置取消事件
        # 取消检查在 (row_idx + 1) % 50 == 0 或 row_idx == total_records - 1 时触发
        call_count = [0]
        def cancel_on_execute(*args, **kwargs):
            call_count[0] += 1
            if call_count[0] >= 1:
                self.updater.cancel_event.set()
            return None

        mock_cursor.execute.side_effect = cancel_on_execute

        with self.assertRaises(CancelledError):
            self.updater.execute_multi_column_update(
                "APPS", "APPS", "EMPLOYEE", "EMP_ID", ["COL1"]
            )


class TestDataUpdaterReportProgress(unittest.TestCase):
    """测试 _report_progress 方法"""

    def setUp(self):
        self.mock_db = MagicMock()
        self.mock_log = MagicMock()
        self.updater = DataUpdater(self.mock_db, self.mock_log)

    def test_report_progress_with_callback(self):
        """有回调时正确调用"""
        callback_calls = []

        def callback(current, total, pct, op, eta=None):
            callback_calls.append((current, total, pct, op))

        self.updater.set_progress_callback(callback)
        self.updater._report_progress(50, 100, "更新数据")

        self.assertEqual(len(callback_calls), 1)
        self.assertEqual(callback_calls[0][0], 50)
        self.assertEqual(callback_calls[0][1], 100)
        self.assertEqual(callback_calls[0][2], 50)
        self.assertEqual(callback_calls[0][3], "更新数据")

    def test_report_progress_without_callback(self):
        """无回调时不抛异常"""
        self.updater.progress_callback = None
        # 不应抛异常
        self.updater._report_progress(50, 100, "test")

    def test_report_progress_backward_compatible_callback(self):
        """向后兼容不接受 eta 参数的回调"""
        callback_calls = []

        def old_callback(current, total, pct, op):
            callback_calls.append((current, total, pct, op))

        self.updater.set_progress_callback(old_callback)
        self.updater._report_progress(25, 100, "导入")

        self.assertEqual(len(callback_calls), 1)

    def test_report_progress_zero_total(self):
        """total=0 时百分比为 0"""
        callback_calls = []

        def callback(current, total, pct, op, eta=None):
            callback_calls.append((current, total, pct, op))

        self.updater.set_progress_callback(callback)
        self.updater._report_progress(0, 0, "test")
        self.assertEqual(callback_calls[0][2], 0)


class TestMapOracleType(unittest.TestCase):
    """测试 _map_oracle_type_to_temp_type 静态方法"""

    def test_varchar2(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("VARCHAR2", 100), "VARCHAR2(100)")

    def test_number(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("NUMBER", 22), "NUMBER")

    def test_date(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("DATE", 7), "DATE")

    def test_clob(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("CLOB", 0), "CLOB")

    def test_unknown_type_defaults_to_varchar2(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("UNKNOWN", 0), "VARCHAR2(4000)")

    def test_none_data_type(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type(None, 0), "VARCHAR2(4000)")

    def test_zero_length_varchar2(self):
        """VARCHAR2 长度为 0 时使用默认 4000"""
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("VARCHAR2", 0), "VARCHAR2(4000)")

    def test_timestamp(self):
        self.assertEqual(DataUpdater._map_oracle_type_to_temp_type("TIMESTAMP", 11), "DATE")


# ============================================================================
# 3. config_manager.py — ConfigManager
# ============================================================================

class TestConfigManagerConnections(unittest.TestCase):
    """测试 add_connection / get_connection / list_connections"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        # Patch ConfigManager 使其使用临时目录
        self.orig_init = ConfigManager.__init__

        def patched_init(self_cm):
            self_cm.app_dir = Path(self.tmpdir)
            self_cm.config_file = Path(self.tmpdir) / "config.json"
            self_cm.config = self_cm.load_config()

        ConfigManager.__init__ = patched_init
        self.cm = ConfigManager()

    def tearDown(self):
        ConfigManager.__init__ = self.orig_init
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_add_and_get_connection(self):
        """添加并获取连接"""
        conn_info = {
            "name": "test_conn",
            "host": "localhost",
            "port": 1521,
            "service": "ORCL",
            "username": "user1",
            "password": "pass1"
        }
        self.cm.add_connection(conn_info)

        # get_connections 返回解密后的密码
        conns = self.cm.get_connections()
        self.assertEqual(len(conns), 1)
        self.assertEqual(conns[0]["name"], "test_conn")
        self.assertEqual(conns[0]["password"], "pass1")

    def test_add_connection_updates_existing(self):
        """同名连接被更新而非追加"""
        conn1 = {"name": "c1", "host": "h1", "port": 1, "service": "s", "username": "u", "password": "p"}
        conn2 = {"name": "c1", "host": "h2", "port": 2, "service": "s", "username": "u", "password": "p2"}
        self.cm.add_connection(conn1)
        self.cm.add_connection(conn2)

        conns = self.cm.get_connections()
        self.assertEqual(len(conns), 1)
        self.assertEqual(conns[0]["host"], "h2")

    def test_get_connection_by_name(self):
        """按名称获取连接"""
        conn = {"name": "myconn", "host": "h", "port": 1, "service": "s", "username": "u", "password": "p"}
        self.cm.add_connection(conn)

        found = self.cm.get_connection_by_name("myconn")
        self.assertIsNotNone(found)
        self.assertEqual(found["name"], "myconn")

        not_found = self.cm.get_connection_by_name("nonexistent")
        self.assertIsNone(not_found)

    def test_delete_connection(self):
        """删除连接"""
        conn = {"name": "del_me", "host": "h", "port": 1, "service": "s", "username": "u", "password": "p"}
        self.cm.add_connection(conn)
        self.assertEqual(len(self.cm.get_connections()), 1)

        self.cm.delete_connection("del_me")
        self.assertEqual(len(self.cm.get_connections()), 0)

    def test_list_connections_empty(self):
        """初始无连接"""
        self.assertEqual(len(self.cm.get_connections()), 0)


class TestConfigManagerSaveLoad(unittest.TestCase):
    """测试 save_config / load_config"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.orig_init = ConfigManager.__init__

        def patched_init(self_cm):
            self_cm.app_dir = Path(self.tmpdir)
            self_cm.config_file = Path(self.tmpdir) / "config.json"
            self_cm.config = self_cm.load_config()

        ConfigManager.__init__ = patched_init

    def tearDown(self):
        ConfigManager.__init__ = self.orig_init
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_save_and_load(self):
        """保存后加载保持一致"""
        cm = ConfigManager()
        cm.config["connections"] = [{"name": "c1", "host": "h"}]
        cm.save_config()

        cm2 = ConfigManager()
        self.assertEqual(len(cm2.config["connections"]), 1)
        self.assertEqual(cm2.config["connections"][0]["name"], "c1")

    def test_save_creates_file(self):
        """save_config 创建文件"""
        cm = ConfigManager()
        cm.save_config()
        self.assertTrue(cm.config_file.exists())

    def test_load_corrupted_file_returns_default(self):
        """损坏的配置文件返回默认配置"""
        cm = ConfigManager()
        # 写入无效 JSON
        with open(cm.config_file, 'w') as f:
            f.write("not valid json{{{")

        loaded = cm.load_config()
        self.assertIn("connections", loaded)
        self.assertIn("templates", loaded)

    def test_default_config_structure(self):
        """默认配置包含必要字段"""
        cm = ConfigManager()
        default = cm.get_default_config()
        self.assertIn("last_used", default)
        self.assertIn("connections", default)
        self.assertIn("templates", default)
        self.assertIn("schema_values", default)


class TestConfigManagerTemplates(unittest.TestCase):
    """测试模板功能"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.orig_init = ConfigManager.__init__

        def patched_init(self_cm):
            self_cm.app_dir = Path(self.tmpdir)
            self_cm.config_file = Path(self.tmpdir) / "config.json"
            self_cm.config = self_cm.load_config()

        ConfigManager.__init__ = patched_init
        self.cm = ConfigManager()

    def tearDown(self):
        ConfigManager.__init__ = self.orig_init
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_add_and_get_template(self):
        """添加并获取模板"""
        ok, msg = self.cm.add_template({"name": "tpl1", "description": "test"})
        self.assertTrue(ok)
        self.assertIn("添加", msg)

        tpl = self.cm.get_template_by_name("tpl1")
        self.assertIsNotNone(tpl)
        self.assertEqual(tpl["description"], "test")

    def test_add_template_update_existing(self):
        """同名模板被更新"""
        self.cm.add_template({"name": "tpl1", "description": "v1"})
        ok, msg = self.cm.add_template({"name": "tpl1", "description": "v2"})
        self.assertTrue(ok)
        self.assertIn("更新", msg)

        tpl = self.cm.get_template_by_name("tpl1")
        self.assertEqual(tpl["description"], "v2")

    def test_add_template_empty_name(self):
        """空名称被拒绝"""
        ok, msg = self.cm.add_template({"name": "", "description": "test"})
        self.assertFalse(ok)
        self.assertIn("不能为空", msg)

    def test_add_template_name_too_long(self):
        """超长名称被拒绝"""
        ok, msg = self.cm.add_template({"name": "a" * 101, "description": "test"})
        self.assertFalse(ok)
        self.assertIn("不能超过100", msg)

    def test_list_templates(self):
        """列出所有模板"""
        self.cm.add_template({"name": "t1"})
        self.cm.add_template({"name": "t2"})
        templates = self.cm.get_templates()
        self.assertEqual(len(templates), 2)

    def test_delete_template(self):
        """删除模板"""
        self.cm.add_template({"name": "del_tpl"})
        self.cm.delete_template("del_tpl")
        self.assertIsNone(self.cm.get_template_by_name("del_tpl"))

    def test_create_template_from_current(self):
        """从当前设置创建模板"""
        ok, msg = self.cm.create_template_from_current(
            name="scenario1",
            description="desc",
            connection_name="conn1",
            target_table="EMPLOYEE",
            key_column="EMP_ID",
            update_columns=["NAME", "AGE"],
            schema="APPS",
            temp_schema="APPS"
        )
        self.assertTrue(ok)

        tpl = self.cm.get_template_by_name("scenario1")
        self.assertIsNotNone(tpl)
        self.assertEqual(tpl["target_table"], "EMPLOYEE")
        self.assertEqual(tpl["update_columns"], ["NAME", "AGE"])
        self.assertIn("created_at", tpl)


class TestConfigManagerPasswordEncryption(unittest.TestCase):
    """测试密码加密存储"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()
        self.orig_init = ConfigManager.__init__

        def patched_init(self_cm):
            self_cm.app_dir = Path(self.tmpdir)
            self_cm.config_file = Path(self.tmpdir) / "config.json"
            self_cm.config = self_cm.load_config()

        ConfigManager.__init__ = patched_init
        self.cm = ConfigManager()

    def tearDown(self):
        ConfigManager.__init__ = self.orig_init
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_password_encrypt_decrypt_roundtrip(self):
        """加密解密往返一致"""
        original = "MyS3cretP@ss!"
        encrypted = password_encrypt(original)
        self.assertTrue(encrypted.startswith("enc:"))
        decrypted = password_decrypt(encrypted)
        self.assertEqual(decrypted, original)

    def test_password_stored_encrypted(self):
        """连接配置中密码以加密形式存储"""
        conn = {"name": "c1", "host": "h", "port": 1, "service": "s", "username": "u", "password": "plain_pass"}
        self.cm.add_connection(conn)

        # 直接读取配置文件，验证密码是加密的
        with open(self.cm.config_file, 'r') as f:
            raw = json.load(f)
        stored_pwd = raw["connections"][0]["password"]
        self.assertTrue(stored_pwd.startswith("enc:"))
        self.assertNotEqual(stored_pwd, "plain_pass")

    def test_password_decrypt_backward_compat_base64(self):
        """兼容旧版 Base64 格式"""
        import base64
        original = "old_password"
        b64_encoded = base64.b64encode(original.encode()).decode()
        decrypted = password_decrypt(b64_encoded)
        self.assertEqual(decrypted, original)

    def test_password_decrypt_plain_text_fallback(self):
        """极旧版明文密码直接返回"""
        # 非 enc: 前缀且非有效 Base64 的字符串
        decrypted = password_decrypt("plaintext_password_!@#$%")
        # 可能是 Base64 解码失败后返回原文
        # 只要不抛异常即可
        self.assertIsInstance(decrypted, str)

    def test_password_decrypt_empty_string(self):
        """空字符串返回空"""
        self.assertEqual(password_decrypt(""), "")

    def test_encode_decode_methods(self):
        """ConfigManager._encode_password / _decode_password 方法"""
        original = "test123"
        encoded = ConfigManager._encode_password(original)
        decoded = ConfigManager._decode_password(encoded)
        self.assertEqual(decoded, original)


# ============================================================================
# 4. excel_handler.py — ExcelHandler
# ============================================================================

def _create_test_excel(file_path, headers, rows):
    """辅助函数：创建测试 Excel 文件"""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(file_path)


class TestExcelHandlerReadExcel(unittest.TestCase):
    """测试 read_excel (validate_excel_structure / validate_multi_column_structure)"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_read_valid_excel(self):
        """读取有效 Excel 文件"""
        fp = os.path.join(self.tmpdir, "test.xlsx")
        _create_test_excel(fp, ["ID", "Name"], [[1, "Alice"], [2, "Bob"]])

        ok, msg, data = ExcelHandler.validate_excel_structure(fp)
        self.assertTrue(ok)
        self.assertEqual(len(data), 2)
        self.assertEqual(data[0]["key_value"], 1)
        self.assertEqual(data[0]["update_value"], "Alice")

    def test_read_nonexistent_file(self):
        """文件不存在"""
        ok, msg, data = ExcelHandler.validate_excel_structure("/nonexistent/file.xlsx")
        self.assertFalse(ok)
        self.assertIn("不存在", msg)

    def test_read_excel_header_only(self):
        """只有表头没有数据"""
        fp = os.path.join(self.tmpdir, "empty.xlsx")
        _create_test_excel(fp, ["ID", "Name"], [])

        ok, msg, data = ExcelHandler.validate_excel_structure(fp)
        self.assertFalse(ok)
        self.assertIn("至少需要包含表头和1行数据", msg)

    def test_read_multi_column_excel(self):
        """读取多列 Excel"""
        fp = os.path.join(self.tmpdir, "multi.xlsx")
        _create_test_excel(fp, ["ID", "Name", "Age"], [[1, "Alice", 25], [2, "Bob", 30]])

        ok, msg, data = ExcelHandler.validate_multi_column_structure(fp, ["Name", "Age"])
        self.assertTrue(ok)
        self.assertEqual(len(data), 2)
        self.assertEqual(data[0]["Name"], "Alice")
        self.assertEqual(data[0]["Age"], 25)

    def test_read_multi_column_insufficient_columns(self):
        """多列 Excel 列数不足"""
        fp = os.path.join(self.tmpdir, "few_cols.xlsx")
        _create_test_excel(fp, ["ID", "Name"], [[1, "Alice"]])

        ok, msg, data = ExcelHandler.validate_multi_column_structure(fp, ["Name", "Age", "Dept"])
        self.assertFalse(ok)
        self.assertIn("至少需要", msg)

    def test_duplicate_keys_detected(self):
        """重复唯一标识值被检测"""
        fp = os.path.join(self.tmpdir, "dup.xlsx")
        _create_test_excel(fp, ["ID", "Name"], [[1, "A"], [1, "B"], [2, "C"]])

        ok, msg, data = ExcelHandler.validate_multi_column_structure(fp, ["Name"])
        self.assertFalse(ok)
        self.assertIn("重复", msg)


class TestExcelHandlerPreviewData(unittest.TestCase):
    """测试 get_preview_data 方法"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_preview_small_file(self):
        """小文件预览"""
        fp = os.path.join(self.tmpdir, "small.xlsx")
        _create_test_excel(fp, ["ID", "Value"], [[i, f"v{i}"] for i in range(10)])

        ok, msg, info = ExcelHandler.get_preview_data(fp)
        self.assertTrue(ok)
        self.assertEqual(info["total_rows"], 10)
        self.assertEqual(len(info["rows"]), 10)
        self.assertFalse(info["has_more"])

    def test_preview_large_file_truncated(self):
        """大文件预览截断到 MAX_PREVIEW_ROWS"""
        fp = os.path.join(self.tmpdir, "large.xlsx")
        _create_test_excel(fp, ["ID", "Value"], [[i, f"v{i}"] for i in range(100)])

        ok, msg, info = ExcelHandler.get_preview_data(fp, max_rows=10)
        self.assertTrue(ok)
        self.assertEqual(info["total_rows"], 100)
        self.assertEqual(len(info["rows"]), 10)
        self.assertTrue(info["has_more"])

    def test_preview_headers(self):
        """预览返回正确的表头"""
        fp = os.path.join(self.tmpdir, "headers.xlsx")
        _create_test_excel(fp, ["EmployeeID", "EmployeeName"], [[1, "Alice"]])

        ok, msg, info = ExcelHandler.get_preview_data(fp)
        self.assertTrue(ok)
        self.assertIn("EmployeeID", info["headers"])
        self.assertIn("EmployeeName", info["headers"])

    def test_preview_nonexistent_file(self):
        """预览不存在的文件"""
        ok, msg, info = ExcelHandler.get_preview_data("/nonexistent.xlsx")
        self.assertFalse(ok)


class TestExcelHandlerValidateData(unittest.TestCase):
    """测试 validate_data (validate_excel_structure + check_duplicate_keys)"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_validate_no_duplicates(self):
        """无重复键验证通过"""
        data_rows = [
            {"row_num": 2, "key_value": 1, "update_value": "A"},
            {"row_num": 3, "key_value": 2, "update_value": "B"},
        ]
        ok, dups = ExcelHandler.check_duplicate_keys(data_rows)
        self.assertTrue(ok)
        self.assertEqual(len(dups), 0)

    def test_validate_with_duplicates(self):
        """有重复键验证失败"""
        data_rows = [
            {"row_num": 2, "key_value": 1, "update_value": "A"},
            {"row_num": 3, "key_value": 1, "update_value": "B"},
            {"row_num": 4, "key_value": 2, "update_value": "C"},
        ]
        ok, dups = ExcelHandler.check_duplicate_keys(data_rows)
        self.assertFalse(ok)
        self.assertIn(1, dups)

    def test_validate_empty_first_column_skipped(self):
        """第一列为空的行被跳过"""
        fp = os.path.join(self.tmpdir, "skip.xlsx")
        _create_test_excel(fp, ["ID", "Name"], [[1, "A"], [None, "B"], [3, "C"]])

        ok, msg, data = ExcelHandler.validate_excel_structure(fp)
        self.assertTrue(ok)
        self.assertEqual(len(data), 2)  # 跳过 None 行


class TestExcelHandlerFileLimits(unittest.TestCase):
    """测试文件大小/行数限制"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_max_file_size_constant(self):
        """MAX_FILE_SIZE 为 10MB"""
        self.assertEqual(MAX_FILE_SIZE, 10 * 1024 * 1024)

    def test_max_rows_constant(self):
        """MAX_ROWS 为 100000"""
        self.assertEqual(MAX_ROWS, 100000)

    def test_max_preview_rows_constant(self):
        """MAX_PREVIEW_ROWS 为 50"""
        self.assertEqual(MAX_PREVIEW_ROWS, 50)

    def test_oversized_file_rejected(self):
        """超大文件被拒绝"""
        fp = os.path.join(self.tmpdir, "big.xlsx")
        _create_test_excel(fp, ["ID", "Name"], [[1, "A"]])

        # Mock os.path.getsize 返回超过限制的大小
        with patch("os.path.getsize", return_value=MAX_FILE_SIZE + 1):
            ok, msg, data = ExcelHandler.validate_excel_structure(fp)
            self.assertFalse(ok)
            self.assertIn("文件大小超过限制", msg)

    def test_too_many_rows_rejected(self):
        """超多行数被拒绝"""
        fp = os.path.join(self.tmpdir, "many_rows.xlsx")
        # 创建一个实际有很多行的文件太慢，使用 mock
        _create_test_excel(fp, ["ID", "Name"], [[1, "A"]])

        from openpyxl import load_workbook
        with patch("src.excel_handler.load_workbook") as mock_load:
            mock_wb = MagicMock()
            mock_sheet = MagicMock()
            mock_sheet.max_row = MAX_ROWS + 2  # header + MAX_ROWS+1 data rows
            mock_sheet.max_column = 2
            mock_sheet.__getitem__ = MagicMock(return_value=[MagicMock(value="ID"), MagicMock(value="Name")])
            mock_wb.active = mock_sheet
            mock_load.return_value = mock_wb

            ok, msg, data = ExcelHandler.validate_excel_structure(fp)
            self.assertFalse(ok)
            self.assertIn("数据行数超过限制", msg)


class TestExcelHandlerExport(unittest.TestCase):
    """测试导出功能"""

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def test_export_to_excel(self):
        """导出数据到 Excel"""
        fp = os.path.join(self.tmpdir, "export.xlsx")
        data = [
            {"key_value": "1", "reason": "test1"},
            {"key_value": "2", "reason": "test2"},
        ]
        ok, msg = ExcelHandler.export_to_excel(data, fp)
        self.assertTrue(ok)
        self.assertTrue(os.path.exists(fp))

    def test_export_empty_data(self):
        """导出空数据"""
        fp = os.path.join(self.tmpdir, "empty_export.xlsx")
        ok, msg = ExcelHandler.export_to_excel([], fp)
        self.assertFalse(ok)
        self.assertIn("没有数据", msg)

    def test_export_failed_records(self):
        """导出失败记录"""
        fp = os.path.join(self.tmpdir, "failed.xlsx")
        records = [
            {"key_value": "1", "update_value": "", "reason": "type mismatch", "timestamp": "2026-01-01"},
        ]
        ok, msg = ExcelHandler.export_failed_records(records, fp)
        self.assertTrue(ok)

    def test_export_failed_records_empty(self):
        """导出空失败记录"""
        fp = os.path.join(self.tmpdir, "no_failed.xlsx")
        ok, msg = ExcelHandler.export_failed_records([], fp)
        self.assertFalse(ok)


# ============================================================================
# 运行所有测试
# ============================================================================

def run_all_tests():
    """运行所有深度核心测试"""
    print("\n" + "=" * 70)
    print("核心模块深度功能测试")
    print("=" * 70)
    print(f"测试时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 70)

    loader = unittest.TestLoader()
    suite = unittest.TestSuite()

    test_classes = [
        # db_connection.py
        TestDBConnectionConnect,
        TestDBConnectionDisconnect,
        TestDBConnectionExecuteSql,
        TestDBConnectionCheckPermissions,
        TestConnectionPool,
        # data_updater.py
        TestDataUpdaterCreateTempTable,
        TestDataUpdaterExecuteMultiColumnUpdate,
        TestDataUpdaterExecuteMergeUpdate,
        TestDataUpdaterCancel,
        TestDataUpdaterReportProgress,
        TestMapOracleType,
        # config_manager.py
        TestConfigManagerConnections,
        TestConfigManagerSaveLoad,
        TestConfigManagerTemplates,
        TestConfigManagerPasswordEncryption,
        # excel_handler.py
        TestExcelHandlerReadExcel,
        TestExcelHandlerPreviewData,
        TestExcelHandlerValidateData,
        TestExcelHandlerFileLimits,
        TestExcelHandlerExport,
    ]

    for test_class in test_classes:
        tests = loader.loadTestsFromTestCase(test_class)
        suite.addTests(tests)

    runner = unittest.TextTestRunner(verbosity=2)
    result = runner.run(suite)

    # 总结
    print("\n" + "=" * 70)
    print("测试总结")
    print("=" * 70)
    print(f"运行测试: {result.testsRun}")
    print(f"成功: {result.testsRun - len(result.failures) - len(result.errors)}")
    print(f"失败: {len(result.failures)}")
    print(f"错误: {len(result.errors)}")
    print("=" * 70)

    # 详细记录每个失败
    if result.failures:
        print("\n--- 失败详情 ---")
        for test, tb in result.failures:
            print(f"\n[FAIL] {test}")
            print(tb)

    if result.errors:
        print("\n--- 错误详情 ---")
        for test, tb in result.errors:
            print(f"\n[ERROR] {test}")
            print(tb)

    if result.wasSuccessful():
        print("\n✓ 所有核心模块深度测试通过！")
    else:
        print("\n✗ 存在测试失败/错误！")

    return result


if __name__ == "__main__":
    result = run_all_tests()
    sys.exit(0 if result.wasSuccessful() else 1)
